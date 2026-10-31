import urllib.request, time, sys, os, collections, re
URL = "https://prefeitura.sp.gov.br/cidade/secretarias/upload/fazenda/arquivos/itbi/GUIAS%20DE%20ITBI%20PAGAS%20%2830092026%29.xlsx"
out = []
t = time.time()
req = urllib.request.Request(URL, headers={"User-Agent": "Mozilla/5.0 LeilaoTrack"})
try:
    with urllib.request.urlopen(req, timeout=180) as r:
        dados = r.read()
    open("/tmp/itbi.xlsx", "wb").write(dados)
    out.append(f"baixou {len(dados)//1024//1024} MB em {time.time()-t:.0f}s")
except Exception as e:
    out.append(f"ERRO download {e}")
    print("::notice title=est::" + " ¦ ".join(out)); sys.exit(0)
import openpyxl
wb = openpyxl.load_workbook("/tmp/itbi.xlsx", read_only=True)
out.append("abas=" + str(wb.sheetnames))
for nome in wb.sheetnames[:3]:
    ws = wb[nome]
    linhas = ws.iter_rows(values_only=True)
    cab = None
    for i, row in enumerate(linhas):
        if cab is None:
            if row and sum(1 for x in row if isinstance(x, str)) >= 5:
                cab = row; out.append(f"[{nome}] cab(linha {i})=" + " | ".join(str(x) for x in row if x is not None))
            continue
        if i < 4 or i in (100, 5000):
            out.append(f"[{nome}] ex: " + " | ".join(str(x)[:30] for x in row))
        if i > 5001: break
    out.append(f"[{nome}] max_row={ws.max_row}")
T = " ¦ ".join(out)
print(T)
for i in range(0, len(T), 3500): print("::notice title=est::" + T[i:i+3500])
