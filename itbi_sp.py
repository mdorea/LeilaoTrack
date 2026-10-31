"""ITBI da cidade de São Paulo: vendas reais (dados abertos da Secretaria da Fazenda).

Baixa as planilhas de guias de ITBI pagas, fica só com compra e venda de casas e
apartamentos inteiros e resume o preço por m² por rua e por bairro. O resumo é
pequeno e fica guardado; cada lote de São Paulo recebe o preço por m² das vendas
reais da mesma rua (ou do bairro).

Fonte: https://prefeitura.sp.gov.br/web/fazenda/w/acesso_a_informacao/31501
Cuidados (da própria base): valor declarado pelo contribuinte; sem rurais nem PPI;
área construída do IPTU (nos apartamentos inclui parte das áreas comuns).
"""
import re
import unicodedata
import urllib.request
from datetime import datetime

PAGINA = "https://prefeitura.sp.gov.br/web/fazenda/w/acesso_a_informacao/31501"
UA = "Mozilla/5.0 (LeilaoTrack; dados abertos ITBI)"

_SEM = re.compile(r"[^A-Z0-9 ]+")
_TIPO_RUA = {"R", "RUA", "AV", "AVENIDA", "AL", "ALAMEDA", "TV", "TRAVESSA", "PC", "PCA", "PRACA", "EST", "ESTRADA",
             "ROD", "RODOVIA", "LGO", "LARGO", "VD", "VIADUTO", "PQ", "PRQ", "VL", "VIELA", "PS", "PASSAGEM", "BC", "BECO"}
_TITULOS = {"SEN", "SENADOR", "GAL", "GEN", "GENERAL", "MAL", "MARECHAL", "DR", "DOUTOR", "DRA", "DOUTORA", "PROF", "PROFESSOR",
            "PROFA", "PROFESSORA", "CEL", "CORONEL", "CAP", "CAPITAO", "PE", "PADRE", "ENG", "ENGENHEIRO", "DES",
            "DESEMBARGADOR", "MIN", "MINISTRO", "PRES", "PRESIDENTE", "VER", "VEREADOR", "DEP", "DEPUTADO", "BRIG",
            "BRIGADEIRO", "COMEND", "COMENDADOR", "TEN", "TENENTE", "SGT", "SARGENTO", "ALM", "ALMIRANTE", "CONS",
            "CONSELHEIRO", "D", "DONA", "DOM", "FREI", "MONS", "MONSENHOR", "ARQ", "ARQUITETO", "MAJ", "MAJOR", "GOV",
            "GOVERNADOR", "PREF", "PREFEITO", "JORN", "JORNALISTA", "MAESTRO", "MTO", "VISC", "VISCONDE", "BR", "BARAO"}
_PARTICULAS = {"DE", "DA", "DO", "DOS", "DAS", "E"}
_ABREV_BAIRRO = {"JD": "JARDIM", "JDM": "JARDIM", "VL": "VILA", "PQ": "PARQUE", "PRQ": "PARQUE", "CJ": "CONJUNTO",
                 "CID": "CIDADE", "STA": "SANTA", "STO": "SANTO", "S": "SAO", "N": "NOSSA", "SRA": "SENHORA",
                 "CHAC": "CHACARA", "RES": "RESIDENCIAL", "SIT": "SITIO", "ALT": "ALTO", "LGO": "LARGO", "PTE": "PONTE",
                 "BL": "BELA", "HAB": "HABITACIONAL", "NUC": "NUCLEO", "LOT": "LOTEAMENTO", "INDL": "INDUSTRIAL"}
_ABREV_RUA = {"STA": "SANTA", "STO": "SANTO", "S": "SAO", "N": "NOSSA", "SRA": "SENHORA", "JD": "JARDIM", "VL": "VILA"}


def _limpo(s):
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().upper()
    return _SEM.sub(" ", s).split()


def chave_rua(s):
    """'R SEN OTAVIO MANGABEIRA' e 'Rua Senador Otávio Mangabeira, 71' → 'OTAVIO MANGABEIRA'."""
    s = str(s or "").split(",")[0]
    s = re.sub(r"\s(n[ºo°]|\d).*$", "", s, flags=re.I) if re.search(r"\s\d", s) else s
    p = _limpo(s)
    if p and p[0] in _TIPO_RUA:
        p = p[1:]
    p = [_ABREV_RUA.get(w, w) for w in p]
    p = [w for w in p if w not in _TITULOS and w not in _PARTICULAS and not w.isdigit()]
    return " ".join(p)


def chave_bairro(s):
    p = [_ABREV_BAIRRO.get(w, w) for w in _limpo(s)]
    return " ".join(w for w in p if w not in _PARTICULAS)


def links_planilhas(get):
    """URLs das planilhas .xlsx na ordem da página da Fazenda (o ano atual vem primeiro)."""
    h = get(PAGINA)
    urls = []
    for u in re.findall(r'href="([^"]+\.xlsx)"', h, re.I):
        u = u.replace("&amp;", "&")
        if u.startswith("/"):
            u = "https://prefeitura.sp.gov.br" + u
        if "itbi" in u.lower() and u not in urls:
            urls.append(u.replace(" ", "%20"))
    return urls


def _num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(".", "").replace(",", "."))
    except ValueError:
        return None


def ler_planilha(caminho, meses=18, hoje=None):
    """Vendas de casas e apartamentos inteiros dos últimos `meses` meses."""
    import openpyxl
    hoje = hoje or datetime.now()
    corte = (hoje.year * 12 + hoje.month) - meses
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    vendas = []
    for nome in wb.sheetnames:
        if not re.match(r"^[A-ZÇ]{3}-\d{4}$", nome.strip().upper()):
            continue
        ws = wb[nome]
        cab = None
        for row in ws.iter_rows(values_only=True):
            if cab is None:
                if row and any(isinstance(x, str) and "Logradouro" in x for x in row):
                    cab = {str(x).strip(): i for i, x in enumerate(row) if x is not None}
                    ix = lambda *nomes: next((cab[k] for k in cab for n in nomes if k.lower().startswith(n.lower())), None)
                    I = dict(rua=ix("Nome do Logradouro"), num=ix("Número"), bairro=ix("Bairro"), cep=ix("CEP"),
                             nat=ix("Natureza"), valor=ix("Valor de Transação"), data=ix("Data de Transação"),
                             prop=ix("Proporção Transmitida"), sit=ix("Situação do SQL"), area=ix("Área Construída"),
                             uso=ix("Descrição do uso"), padrao=ix("Padrão (IPTU)"))
                continue
            try:
                if not str(row[I["nat"]] or "").startswith("1."):          # só compra e venda
                    continue
                if "predial" not in str(row[I["sit"]] or "").lower():
                    continue
                if (_num(row[I["prop"]]) or 0) < 99.9:                       # imóvel inteiro
                    continue
                uso = str(row[I["uso"]] or "").upper()
                tipo = "Apartamento" if uso.startswith("APARTAMENTO") else "Casa" if uso.startswith("RESID") and "COLETIVA" not in uso else None
                if not tipo:
                    continue
                area, valor = _num(row[I["area"]]), _num(row[I["valor"]])
                if not area or not valor or not (20 <= area <= 1000):
                    continue
                m2 = valor / area
                if not (500 <= m2 <= 100000):
                    continue
                d = row[I["data"]]
                if isinstance(d, datetime) and d.year * 12 + d.month < corte:
                    continue
                vendas.append((tipo, chave_rua(row[I["rua"]]), chave_bairro(row[I["bairro"]]),
                               str(row[I["cep"]] or "").zfill(8)[:5], m2))
            except (IndexError, TypeError):
                continue
    return vendas


def _q(ordem, p):
    i = (len(ordem) - 1) * p
    a = int(i)
    return ordem[a] + (ordem[min(a + 1, len(ordem) - 1)] - ordem[a]) * (i - a)


def resumir(vendas, min_rua=5, min_bairro=12):
    """{'r': {tipo|rua: [n, p25, mediana, p75]}, 'b': {tipo|bairro: [...]}, 'c': {tipo|cep5: [...]}}"""
    grupos = {"r": {}, "b": {}, "c": {}}
    # bairro em branco em parte das guias: usa o bairro mais comum do mesmo CEP (5 dígitos)
    cep_bairro = {}
    for t, r, b, c, m in vendas:
        if b:
            cep_bairro.setdefault(c, {}).setdefault(b, 0)
            cep_bairro[c][b] += 1
    cep_bairro = {c: max(d, key=d.get) for c, d in cep_bairro.items()}
    for t, r, b, c, m in vendas:
        b = b or cep_bairro.get(c, "")
        if r:
            grupos["r"].setdefault(t + "|" + r, []).append(m)
        if b:
            grupos["b"].setdefault(t + "|" + b, []).append(m)
        grupos["c"].setdefault(t + "|" + c, []).append(m)
    res = {}
    for k, mins in (("r", min_rua), ("b", min_bairro), ("c", min_bairro)):
        res[k] = {}
        for g, vals in grupos[k].items():
            if len(vals) < mins:
                continue
            o = sorted(vals)
            res[k][g] = [len(o), round(_q(o, .25)), round(_q(o, .5)), round(_q(o, .75))]
    return res


def para_lote(l, resumo):
    """Preço por m² das vendas reais perto do lote: mesma rua (melhor) ou mesmo bairro."""
    if not resumo or l.get("uf") != "SP" or chave_bairro(l.get("cidade")) != "SAO PAULO":
        return None
    t = l.get("tipo")
    if t not in ("Casa", "Apartamento"):
        return None
    r = chave_rua(l.get("endereco"))
    if r and len(r) >= 4 and (t + "|" + r) in resumo.get("r", {}):
        n, p25, med, p75 = resumo["r"][t + "|" + r]
        return {"nivel": "rua", "n": n, "p25": p25, "m2": med, "p75": p75}
    b = chave_bairro(l.get("bairro"))
    if b and (t + "|" + b) in resumo.get("b", {}):
        n, p25, med, p75 = resumo["b"][t + "|" + b]
        return {"nivel": "bairro", "n": n, "p25": p25, "m2": med, "p75": p75}
    return None
